"""Leitura de arquivos tabulares (CSV e planilha XLSX)."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from gestlog.ingestion.erros import ArquivoVazio, FormatoArquivoInvalido

_EXTENSOES_CSV = (".csv", ".txt")
_EXTENSOES_XLSX = (".xlsx", ".xlsm")


def ler_tabela(
    conteudo: bytes, nome_arquivo: str = ""
) -> tuple[tuple[str, ...], tuple[tuple[int, dict[str, str]], ...]]:
    """Lê ``conteudo`` e devolve o cabeçalho e as linhas numeradas."""
    if not conteudo:
        raise ArquivoVazio()
    formato = _detectar_formato(nome_arquivo, conteudo)
    bruto = _ler_xlsx(conteudo) if formato == "xlsx" else _ler_csv(conteudo)
    return _montar(bruto)


def _detectar_formato(nome_arquivo: str, conteudo: bytes) -> str:
    nome = nome_arquivo.lower()
    if nome.endswith(_EXTENSOES_XLSX) or conteudo[:4] == b"PK\x03\x04":
        return "xlsx"
    if nome.endswith(_EXTENSOES_CSV) or "." not in nome:
        return "csv"
    raise FormatoArquivoInvalido(nome_arquivo)


def _ler_csv(conteudo: bytes) -> list[list[str]]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("latin-1")
    delimitador = _detectar_delimitador(texto)
    return [
        list(linha) for linha in csv.reader(io.StringIO(texto), delimiter=delimitador)
    ]


def _detectar_delimitador(texto: str) -> str:
    primeira = texto.splitlines()[0] if texto.splitlines() else ""
    return ";" if primeira.count(";") > primeira.count(",") else ","


def _ler_xlsx(conteudo: bytes) -> list[list[str]]:
    planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    ativa = planilha.active
    return [
        [_celula(valor) for valor in linha]
        for linha in ativa.iter_rows(values_only=True)
    ]


def _celula(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "sim" if valor else "não"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def _montar(
    bruto: list[list[str]],
) -> tuple[tuple[str, ...], tuple[tuple[int, dict[str, str]], ...]]:
    indice = next(
        (i for i, linha in enumerate(bruto) if any(c.strip() for c in linha)), None
    )
    if indice is None:
        raise ArquivoVazio()
    cabecalho = tuple(celula.strip().lower() for celula in bruto[indice])
    linhas: list[tuple[int, dict[str, str]]] = []
    for deslocamento, linha in enumerate(bruto[indice + 1 :]):
        if not any(celula.strip() for celula in linha):
            continue
        numero = indice + 2 + deslocamento
        registro = {
            cabecalho[i]: (linha[i] if i < len(linha) else "")
            for i in range(len(cabecalho))
        }
        linhas.append((numero, registro))
    return cabecalho, tuple(linhas)
